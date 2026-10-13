"""Export docs/Chapter-4-Results-and-Discussion.md to a Word manuscript with embedded figures.

Usage: python scripts/export_chapter4_docx.py
Output: docs/Chapter_4_Results_and_Discussion_Revised.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "Chapter-4-Results-and-Discussion.md"
OUT = ROOT / "docs" / "Chapter_4_Results_and_Discussion_Revised.docx"

LATEX = [
    (r"\\text\{([^}]*)\}", r"\1"),
    (r"\\pm", "±"),
    (r"\\le\b", "≤"),
    (r"\\ge\b", "≥"),
    (r"\\times", "×"),
    (r"\\alpha", "α"),
    (r"\\mu", "μ"),
    (r"\\sigma", "σ"),
    (r"\\,", " "),
    (r"_\{([^}]*)\}", r"_\1"),
    (r"\\", ""),
]


def clean_math(text: str) -> str:
    def repl(match):
        body = match.group(1)
        for pattern, sub in LATEX:
            body = re.sub(pattern, sub, body)
        return body.strip()

    return re.sub(r"\$([^$]+)\$", repl, text)


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


def add_inline(paragraph, text: str, size=None, italic_all=False):
    text = clean_math(text)
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if italic_all:
            run.italic = True
        if size:
            run.font.size = Pt(size)


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    props.append(shd)


def add_table(doc, rows):
    header, body = rows[0], rows[2:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, text in enumerate(header):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        add_inline(cell.paragraphs[0], text, size=8.5)
        for run in cell.paragraphs[0].runs:
            run.bold = True
        shade(cell, "D9E2F3")
    for row in body:
        cells = table.add_row().cells
        for idx in range(len(header)):
            cells[idx].text = ""
            add_inline(cells[idx].paragraphs[0], row[idx] if idx < len(row) else "", size=8.5)
    doc.add_paragraph()


def split_row(line: str):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(1)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped or stripped == "---":
            i += 1
            continue

        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(split_row(lines[i]))
                i += 1
            add_table(doc, block)
            continue

        image = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image:
            path = (SRC.parent / image.group(2)).resolve()
            if path.exists():
                doc.add_picture(str(path), width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_paragraph(f"[Missing figure file: {image.group(2)}]")
            i += 1
            continue

        if stripped.startswith("$$"):
            body = stripped.strip("$").strip()
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(para, f"${body}$")
            i += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if heading:
            level = min(len(heading.group(1)), 4)
            text = clean_math(heading.group(2))
            doc.add_heading(re.sub(r"[*`]", "", text), level=level if level > 1 else 1)
            i += 1
            continue

        if stripped.startswith(">"):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.4)
            add_inline(para, stripped.lstrip("> ").strip(), size=10)
            i += 1
            continue

        if re.match(r"^(\d+)\.\s+", stripped):
            para = doc.add_paragraph(style="List Number")
            add_inline(para, re.sub(r"^\d+\.\s+", "", stripped))
            i += 1
            continue

        if stripped.startswith("- "):
            para = doc.add_paragraph(style="List Bullet")
            add_inline(para, stripped[2:])
            i += 1
            continue

        if re.match(r"^\*\*.+\*\*\s{2,}$", line) or line.endswith("  "):
            para = doc.add_paragraph()
            add_inline(para, stripped)
            para.paragraph_format.space_after = Pt(0)
            i += 1
            continue

        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(8)
        add_inline(para, stripped)
        i += 1

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
