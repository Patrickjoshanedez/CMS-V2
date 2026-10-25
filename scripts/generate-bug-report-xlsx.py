"""Generate docs/CMS-V2-Bug-Report-Log.xlsx (bug tracker with snapshot support).

Run: uv run --with openpyxl scripts/generate-bug-report-xlsx.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

OUT = Path(__file__).resolve().parent.parent / "docs" / "CMS-V2-Bug-Report-Log.xlsx"
OUT.parent.mkdir(exist_ok=True)
(OUT.parent / "bug-snapshots").mkdir(exist_ok=True)

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

SEVERITY = ["Critical", "High", "Medium", "Low"]
STATUS = ["Open", "Triaged", "In Progress", "Fixed", "Verified", "Won't Fix", "Duplicate"]
PHASE = ["Phase 0", "Phase 1", "Phase 2", "Phase 3", "Phase 4", "N/A"]
ENV = ["Local dev", "Docker", "LAN deploy", "Production"]
ROLE = ["student", "instructor", "faculty", "admin"]
FREQ = ["Always", "Often", "Sometimes", "Once"]
THEME = ["Light", "Dark", "Both"]
YESNO = ["Yes", "No"]

# (header, width, validation list or None)
COLUMNS = [
    ("Bug ID", 11, None),
    ("Title", 38, None),
    ("Status", 13, STATUS),
    ("Severity", 11, SEVERITY),
    ("Date Reported", 14, None),
    ("Reported By", 18, None),
    ("Reporter Role", 14, ROLE),
    ("Assigned To", 16, None),
    ("Capstone Phase", 14, PHASE),
    ("Area / Module", 20, None),
    ("Page URL", 34, None),
    ("Environment", 14, ENV),
    ("App Version / Commit", 18, None),
    ("Browser + Version", 18, None),
    ("OS / Device", 18, None),
    ("Viewport (e.g. 1440x900)", 16, None),
    ("Theme", 10, THEME),
    ("Description", 45, None),
    ("Steps to Reproduce", 50, None),
    ("Expected Result", 36, None),
    ("Actual Result (exact error text)", 36, None),
    ("Frequency", 12, FREQ),
    ("Snapshot 1 (file in docs/bug-snapshots)", 34, None),
    ("Snapshot 2", 30, None),
    ("Snapshot 3", 30, None),
    ("Console Error", 40, None),
    ("Network Request (method, URL, status, body)", 40, None),
    ("Server Log", 40, None),
    ("Related IDs (project/team/chapter/defense)", 30, None),
    ("Record State (e.g. workflowPhase=2)", 28, None),
    ("Suspected Cause / Files", 34, None),
    ("Workaround", 30, None),
    ("Impact", 30, None),
    ("Root Cause (dev)", 34, None),
    ("Fix Summary (dev)", 34, None),
    ("Files Changed (dev)", 30, None),
    ("Commit / PR (dev)", 20, None),
    ("Test Added (dev)", 24, YESNO),
    ("Verified By / Date (dev)", 22, None),
    ("Regression Checked light/dark desktop/mobile (dev)", 20, YESNO),
]

wb = Workbook()

# ---- Bug Log sheet ----
ws = wb.active
ws.title = "Bug Log"
for i, (h, w, _) in enumerate(COLUMNS, 1):
    c = ws.cell(row=1, column=i, value=h)
    c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 48
ws.freeze_panes = "C2"
MAX_ROW = 500
ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{MAX_ROW}"

# Example row (delete after reading)
example = {
    "Bug ID": "BUG-001",
    "Title": "EXAMPLE: Chapter 2 upload shows blank preview in dark mode",
    "Status": "Open", "Severity": "Medium",
    "Reporter Role": "student", "Capstone Phase": "Phase 2",
    "Area / Module": "Chapter Review",
    "Page URL": "http://localhost:5173/projects/<id>/chapters/2",
    "Environment": "Docker", "Viewport (e.g. 1440x900)": "1440x900", "Theme": "Dark",
    "Description": "The document viewer renders a white blank area after upload.",
    "Steps to Reproduce": "1. Log in as student\n2. Open Chapter 2\n3. Upload a .docx\n4. Toggle dark mode",
    "Expected Result": "Document renders using SophisticatedDocumentViewer.",
    "Actual Result (exact error text)": "Blank panel, no error shown.",
    "Frequency": "Always",
    "Snapshot 1 (file in docs/bug-snapshots)": "BUG-001_01_dark_desktop.png",
}
for i, (h, _, _) in enumerate(COLUMNS, 1):
    c = ws.cell(row=2, column=i, value=example.get(h))
    c.font = Font(italic=True, color="7F7F7F")

for r in range(2, MAX_ROW + 1):
    for i in range(1, len(COLUMNS) + 1):
        c = ws.cell(row=r, column=i)
        c.alignment, c.border = WRAP, BORDER
    ws.cell(row=r, column=5).number_format = "yyyy-mm-dd"

for i, (_, _, opts) in enumerate(COLUMNS, 1):
    if opts:
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        col = get_column_letter(i)
        dv.add(f"{col}2:{col}{MAX_ROW}")

sev = f"D2:D{MAX_ROW}"
for val, color in (("Critical", "F8696B"), ("High", "FFB366"), ("Medium", "FFE699"), ("Low", "C6E0B4")):
    ws.conditional_formatting.add(
        sev, FormulaRule(formula=[f'$D2="{val}"'], fill=PatternFill("solid", bgColor=color, fgColor=color)))
for val, color in (("Verified", "C6E0B4"), ("Fixed", "DDEBF7"), ("Open", "FCE4D6")):
    ws.conditional_formatting.add(
        f"C2:C{MAX_ROW}", FormulaRule(formula=[f'$C2="{val}"'], fill=PatternFill("solid", bgColor=color, fgColor=color)))

# ---- Snapshots sheet (embedded images go here) ----
ss = wb.create_sheet("Snapshots")
for i, (h, w) in enumerate([("Bug ID", 12), ("Snapshot #", 11), ("Caption / What it shows", 50),
                            ("File name (docs/bug-snapshots)", 42), ("Embedded image (Insert > Pictures)", 60)], 1):
    c = ss.cell(row=1, column=i, value=h)
    c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
    ss.column_dimensions[get_column_letter(i)].width = w
ss.freeze_panes = "A2"
ss.append(["BUG-001", 1, "EXAMPLE: blank viewer in dark mode", "BUG-001_01_dark_desktop.png",
           "Paste/insert the picture in this cell, set row height ~220"])
for i in range(1, 6):
    ss.cell(row=2, column=i).font = Font(italic=True, color="7F7F7F")
    ss.cell(row=2, column=i).alignment = WRAP
ss.row_dimensions[2].height = 60

# ---- Instructions sheet ----
ins = wb.create_sheet("Instructions", 0)
ins.column_dimensions["A"].width = 26
ins.column_dimensions["B"].width = 100
rows = [
    ("CMS-V2 Bug Report Log", ""),
    ("", ""),
    ("How to file", "1. Go to the 'Bug Log' sheet and add one row per bug (row 2 is an example; overwrite or delete it)."),
    ("", "2. Use the next Bug ID (BUG-002, BUG-003, ...)."),
    ("", "3. Fill dropdown columns (Status, Severity, Role, Phase, Environment, Theme, Frequency)."),
    ("", "4. Save screenshots to docs/bug-snapshots/ named BUG-<ID>_<step>_<light|dark>_<desktop|mobile>.png and list the file name in the Snapshot columns."),
    ("", "5. Optionally embed the picture on the 'Snapshots' sheet (Insert > Pictures > Place in Cell) so the workbook is self-contained."),
    ("", "6. Include the exact error text, browser console output, and the failing network request (Network tab) whenever there is an error."),
    ("", "7. Columns marked (dev) are filled in by the developer who resolves the bug."),
    ("", ""),
    ("Severity guide", "Critical: data loss, security issue, or a phase gate blocked for everyone."),
    ("", "High: core workflow broken, no workaround."),
    ("", "Medium: broken but a workaround exists."),
    ("", "Low: cosmetic, typo, minor layout glitch."),
    ("", ""),
    ("Rules", "Never paste passwords, tokens, or .env values. Blur personal data (student IDs, emails) in snapshots."),
    ("", "For UI bugs, capture light + dark mode at 1440x900 desktop and 390x844 mobile when possible."),
]
for r in rows:
    ins.append(list(r))
ins["A1"].font = Font(bold=True, size=16, color="1F3864")
for r in range(1, len(rows) + 1):
    ins.cell(row=r, column=1).font = Font(bold=True, size=16 if r == 1 else 11,
                                          color="1F3864" if r == 1 else "000000")
    ins.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")

wb.save(OUT)
print(f"Saved {OUT}")
